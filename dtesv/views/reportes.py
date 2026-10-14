from django.http import HttpResponse
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment
from openpyxl.utils import get_column_letter
from datetime import date, datetime
from django.shortcuts import render
from django.db.models import Count, Sum
from dtesv.models import Documentos
from django.views.decorators.csrf import csrf_exempt
from django.http import JsonResponse

import io
import logging

logger = logging.getLogger(__name__)

# Funcion para exportar registros a Excel
def exportar_resumen_clientes_excel(request):
    desde = request.GET.get('desde')
    hasta = request.GET.get('hasta')
    tipo = request.GET.get('tipo')
    empresa = request.GET.get('empresa')
    logger.error (f"El valor de empresa es: {empresa}")
    if not empresa:
        empresa = 'Empresa por Defecto'

    if not desde or not hasta:
        return HttpResponse("Parámetros 'desde' y 'hasta' son requeridos.", status=400)

    wb = Workbook()
    ws = wb.active
    ws.title = "Resumen Cliente-Facturas"

    # Fuente negrita para encabezados
    bold_font = Font(bold=True, size=12)
    center_align = Alignment(horizontal='center')

    # --- Encabezados superiores ---
    ws.merge_cells('A1:H1')
    ws['A1'] = empresa
    ws['A1'].font = Font(bold=True, size=14)
    ws['A1'].alignment = center_align

    ws.merge_cells('A2:H2')
    ws['A2'] = "Reporte Facturas por Cliente"  # O nombre dinámico según tipo u otro parámetro
    ws['A2'].font = Font(bold=True, size=13)
    ws['A2'].alignment = center_align

    ws.merge_cells('A3:H3')
    ws['A3'] = f"Fecha de generación: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}"
    ws['A3'].alignment = center_align

    ws.merge_cells('A4:H4')
    ws['A4'] = f"Fechas consultadas: {desde} || {hasta} "
    ws['A4'].alignment = center_align

    # Línea vacía para separar
    ws.append([])

    # Encabezados
    headers = ['Codigo', 'Nombre', 'Numero Documento','NRC', 'Tipo Receptor', 'Tipo DTE', 'Total Documentos', 'Total Facturas']
    ws.append(headers)
    header_row = ws[ws.max_row]
    for cell in header_row:
        cell.font = bold_font
        cell.alignment = center_align

    # Aqui se recopilan los datos (puedes usar tus modelos aquí)
    data = resumen_por_cliente(desde, hasta, tipo)

    # Cuerpo de datos
    for fila in data:
        tipo_doc = fila.get('tipodocumento')
        total = fila.get('total') or 0

        # Si el DTE es de tipo '05', lo consideras negativo (Nota de Crédito)
        if tipo_doc == '05':
            total = -total        
        ws.append([
            fila['receptor_id__codigo'],
            fila['receptor_id__nombre'],
            fila['receptor_id__numdocumento'],
            fila['receptor_id__nrc'],
            fila['receptor_id__tiporeceptor'],
            tipo_doc,
            fila['cantidad'],
            float(total)
        ])

    # --- Auto-ajustar ancho de columnas ---
    for col_idx, col in enumerate(ws.columns, 1):
        max_length = 0
        col_letter = get_column_letter(col_idx)
        for cell in col:
            if cell.value:
                max_length = max(max_length, len(str(cell.value)))
        if max_length > 0:
            ws.column_dimensions[col_letter].width = max_length + 2

    # Crear respuesta HTTP con archivo
    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    filename = f"Resumen_Clientes.xlsx"
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    # logger.error (f"Respuesta de exportacion: {response}")

    wb.save(response)
    return response



# Funcion auxiliar: resumen agrupado por cliente
def resumen_por_cliente(start_date=None, end_date=None, tipo_receptor=None):
    filtros = {'receptor_id__isnull':False}


    # ✅ Validar fechas
    if not start_date:
        start_date = date.today().isoformat()
    if not end_date:
        end_date = date.today().isoformat()

    filtros['fecEmi__range'] = [start_date, end_date]

    # ✅ Validar tipo receptor
    if tipo_receptor in ['cliente', 'proveedor']:
        filtros['receptor_id__tiporeceptor'] = tipo_receptor


    resumen = (
        Documentos.objects
        .filter(**filtros)
        .values(
            'receptor_id__codigo',
            'receptor_id__nombre',
            'receptor_id__numdocumento',
            'receptor_id__nrc',
            'receptor_id__tiporeceptor',
            'tipodocumento'
        )
        .annotate(
            cantidad = Count('codigoGeneracion'),
            total= Sum('totalPagar')
        )
        .order_by('receptor_id__codigo', 'tipodocumento')
    )
    logger.error(f"Reportes obtuvo los documentos: {resumen}")
    return list(resumen)

# Vista principal que carga el HTML con los datos
def reportes_view(request):
    logger.error("Cargando la vista del template")

    # lógica de backend si se requiere
    start_date = request.GET.get('desde')
    end_date = request.GET.get('hasta')
    tipo = request.GET.get('tipo')

    # Aplicar fecha actual solo si vienen vacías o None
    if not start_date or start_date.strip() == "":
        start_date = date.today().isoformat()
    if not end_date or end_date.strip() == "":
        end_date = date.today().isoformat()

    logger.error(f"Usando desde={start_date} hasta={end_date} tipo={tipo}")

    data = resumen_por_cliente(start_date, end_date, tipo)
    context = {
        'resumen_clientes': data,
        'desde': start_date,
        'hasta': end_date,
        'tipo': tipo,
        'total_clientes': len(data),  # 👈 esto es lo nuevo
    }

    return render(request, 'reportes.html', context)


# @login_required
def exportar_documentos_excel(request):
    desde = request.GET.get('desde')
    hasta = request.GET.get('hasta')
    empresa = request.GET.get('empresa')

    if not desde or not hasta:
        return HttpResponse("Parámetros 'desde' y 'hasta' son requeridos.", status=400)

    documentos = Documentos.objects.filter(fecEmi__range=[desde, hasta])

    wb = Workbook()
    ws = wb.active
    ws.title = "Documentos"

    # Fuente negrita para encabezados
    bold_font = Font(bold=True, size=12)
    center_align = Alignment(horizontal='center')

    # --- Encabezados superiores ---
    ws.merge_cells('A1:M1')
    ws['A1'] = empresa
    ws['A1'].font = Font(bold=True, size=14)
    ws['A1'].alignment = center_align

    ws.merge_cells('A2:N2')
    ws['A2'] = "Reporte Estado de Documentos"  # O nombre dinámico según tipo u otro parámetro
    ws['A2'].font = Font(bold=True, size=13)
    ws['A2'].alignment = center_align

    ws.merge_cells('A3:N3')
    ws['A3'] = f"Fecha de generación: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}"
    ws['A3'].alignment = center_align

    ws.merge_cells('A4:N4')
    ws['A4'] = f"Fechas consultadas: {desde} || {hasta} "
    ws['A4'].alignment = center_align    

    # Línea vacía para separar
    ws.append([])

    headers = [
        'No Interno', 'Tipo Documento', 'Codigo', 'Receptor', 'Estado', 'Codigo de Generación',
        'Sello de Recibido', 'Fecha de Proceso MH', 'Documento Relacionado', 'Total de Documento', 'Fecha Anulación',
        'Fecha Emision', 'Hora Emision', 'Observaciones del Proceso', 'Observaciones de MH'
    ]
    ws.append(headers)
    # Definir formato al encabezado
    header_row = ws[ws.max_row]
    for cell in header_row:
        cell.font = bold_font
        cell.alignment = center_align

    for doc in documentos:
        ws.append([
            doc.num_documento,
            str(doc.tipodocumento),
            (doc.receptor_id.codigo if doc.receptor_id else (doc.proveedor_id.codigo if doc.proveedor_id else '')),
            (doc.receptor_id.nombre if doc.receptor_id else (doc.proveedor_id.nombre if doc.proveedor_id else '')),  # ajustar según FK
            doc.estado,
            doc.codigoGeneracion,
            doc.selloRecibido,
            doc.fecha_proceso_mh.strftime('%Y-%m-%d %H:%M:%S') if doc.fecha_proceso_mh else '',
            doc.numeroDocumento_rel_guid,
            float(doc.totalPagar),
            doc.fecha_anula_mh.strftime('%Y-%m-%d %H:%M:%S') if doc.fecha_anula_mh else '',            
            doc.fecEmi.strftime('%Y-%m-%d'),
            doc.horEmi.strftime('%H:%M:%S') if doc.horEmi else '',
            doc.observacion_proceso or '',
            doc.observaciones_mh or ''
        ])

    # --- Auto-ajustar ancho de columnas ---
    for col_idx, col in enumerate(ws.columns, 1):
        max_length = 0
        col_letter = get_column_letter(col_idx)
        for cell in col:
            if cell.value:
                max_length = max(max_length, len(str(cell.value)))
        if max_length > 0:
            ws.column_dimensions[col_letter].width = max_length + 2

    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    filename = f"Reporte_Documentos_{desde}_al_{hasta}.xlsx"
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    wb.save(response)
    return response
